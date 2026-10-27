"""把 docs/使用说明.md 重建成 docs/使用说明.docx（沿用旧档的字号/配色/表样式）。

为什么留在仓库里：Word 版没有测试守着，2.1.0 → 2.3.0 两版都没重生成（GitHub About 是同一个漏法）。
需要 python-docx：`python docs/tools/build-usage-docx.py`。改 md 之后跑一次，再看 `git status`。

版式取自 v2.1.0 那份 docx：Letter（7772400×10058400 EMU）、四边 0.9in、
标题 22pt 粗 #2F2A45、副标题 11pt #5A5568、版本行 9pt #7A7488、正文 10.5pt、
表用 Table Grid、代码用 Consolas。
"""

from pathlib import Path
import re

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Emu, Pt, RGBColor

REPO = Path(__file__).resolve().parents[2]
SRC = REPO / "docs" / "使用说明.md"
DST = REPO / "docs" / "使用说明.docx"

TITLE = RGBColor(0x2F, 0x2A, 0x45)
SUB = RGBColor(0x5A, 0x55, 0x68)
MUTED = RGBColor(0x7A, 0x74, 0x88)

BODY = 10.5


def build():
    doc = docx.Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Emu(7772400), Emu(10058400)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Emu(822960))

    def runs_inline(par, text, size=BODY, bold=False, mono=False):
        """把 **粗** 与 `代码` 拆成 run，其余按普通文字。"""
        for tok in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
            if not tok:
                continue
            if tok.startswith("**") and tok.endswith("**"):
                r = par.add_run(tok[2:-2])
                r.bold = True
            elif tok.startswith("`") and tok.endswith("`"):
                r = par.add_run(tok[1:-1])
                r.font.name = "Consolas"
            else:
                r = par.add_run(tok)
                if bold:
                    r.bold = True
            if mono:
                r.font.name = "Consolas"
            r.font.size = Pt(9.5 if mono else size)

    def flush_table(rows):
        rows = [r for r in rows if not re.fullmatch(r"\|[-: |]+\|", r.strip())]
        cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
        if not cells:
            return
        width = max(len(r) for r in cells)
        table = doc.add_table(rows=len(cells), cols=width)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        for i, row in enumerate(cells):
            for j in range(width):
                cell = table.cell(i, j)
                cell.paragraphs[0].text = ""
                runs_inline(cell.paragraphs[0], row[j] if j < len(row) else "", size=9.5, bold=(i == 0))
        doc.add_paragraph()

    lines = SRC.read_text(encoding="utf-8").splitlines()
    i = 0
    pending_table = []
    in_code = False
    while i < len(lines):
        line = lines[i]
        if line.startswith("```"):
            in_code = not in_code
            i += 1
            continue
        if in_code:
            par = doc.add_paragraph()
            r = par.add_run(line)
            r.font.name = "Consolas"
            r.font.size = Pt(10)
            i += 1
            continue
        if line.startswith("|"):
            pending_table.append(line)
            i += 1
            continue
        if pending_table:
            flush_table(pending_table)
            pending_table = []
        if not line.strip() or line.strip() == "---":
            i += 1
            continue
        if line.startswith("# "):
            par = doc.add_paragraph()
            r = par.add_run(line[2:])
            r.bold = True
            r.font.size = Pt(22)
            r.font.color.rgb = TITLE
        elif line.startswith("### "):
            par = doc.add_paragraph()
            r = par.add_run(line[4:])
            r.bold = True
            r.font.size = Pt(BODY)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=1)
        elif line.startswith("> "):
            text = line[2:]
            par = doc.add_paragraph()
            runs_inline(par, text, size=9 if text.startswith("版本") else 11)
            for r in par.runs:
                r.font.color.rgb = MUTED if text.startswith("版本") else SUB
        elif line.startswith("  - "):
            runs_inline(doc.add_paragraph(style="List Bullet 2"), line[4:])
        elif line.startswith("- "):
            runs_inline(doc.add_paragraph(style="List Bullet"), line[2:])
        else:
            runs_inline(doc.add_paragraph(), line)
        i += 1
    if pending_table:
        flush_table(pending_table)

    doc.save(DST)
    return DST


if __name__ == "__main__":
    print("saved", build())
